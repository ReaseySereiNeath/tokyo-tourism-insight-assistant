import io

import openpyxl
from openpyxl.styles import Font

from app.importers.jnto import parse_jnto_monthly_workbook
from app.importers.service import run_import

FEEDBACK_HEADER = "text,language,source,rating,permission_basis,publication_date,collection_date\n"


def count(conn, table):
    return conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]


def test_missing_required_column_rejects_file(real_conn):
    batch = run_import(real_conn, "real", "feedback", "fb.csv", b"text,language\nGreat tour,en\n")
    assert batch["status"] == "rejected"
    assert "Missing required column" in batch["errors"][0]["message"]
    assert count(real_conn, "feedback") == 0


def test_invalid_rows_reported_with_row_numbers_and_nothing_imported(real_conn):
    csv = (FEEDBACK_HEADER
           + "Lovely walk,en,Survey,5,Own survey,2025-01-02,2025-02-01\n"
           + ",en,Survey,5,Own survey,,2025-02-01\n"                      # row 3: no text
           + "Good food,en,Survey,5,Own survey,02/01/2025,2025-02-01\n")  # row 4: ambiguous date
    batch = run_import(real_conn, "real", "feedback", "fb.csv", csv.encode())
    assert batch["status"] == "rejected"
    by_row = {e["row"]: e for e in batch["errors"]}
    assert set(by_row) == {3, 4}
    assert by_row[3]["column"] == "text"
    assert "ambiguous" in by_row[4]["message"]
    assert count(real_conn, "feedback") == 0  # all-or-nothing


def test_importing_same_file_twice_does_not_inflate_counts(real_conn):
    csv = (FEEDBACK_HEADER + "Lovely walk,en,Survey,5,Own survey,2025-01-02,2025-02-01\n"
           + "Too fast,en,Survey,3,Own survey,,2025-02-01\n").encode()
    first = run_import(real_conn, "real", "feedback", "fb.csv", csv)
    second = run_import(real_conn, "real", "feedback", "fb_copy.csv", csv)
    assert (first["rows_inserted"], second["rows_inserted"], second["rows_duplicate"]) == (2, 0, 2)
    assert count(real_conn, "feedback") == 2
    assert any("already imported" in w for w in second["warnings"])


def test_duplicate_rows_within_one_file_are_counted_once(real_conn):
    row = "Lovely walk,en,Survey,5,Own survey,2025-01-02,2025-02-01\n"
    batch = run_import(real_conn, "real", "feedback", "fb.csv", (FEEDBACK_HEADER + row + row).encode())
    assert (batch["rows_inserted"], batch["rows_duplicate"]) == (1, 1)


VS_HEADER = "reporting_month,geography,visitor_origin,metric,value,unit,value_status,source,collection_date\n"


def test_revised_statistic_updates_value_and_keeps_history(real_conn):
    run_import(real_conn, "real", "visitor_stats", "a.csv",
               (VS_HEADER + "2025-01,Japan,United States,visitor_arrivals,100,persons,estimate,JNTO,2025-02-01\n").encode())
    batch = run_import(real_conn, "real", "visitor_stats", "b.csv",
                       (VS_HEADER + "2025-01,Japan,United States,visitor_arrivals,120,persons,final,JNTO,2025-03-01\n").encode())
    assert batch["rows_updated"] == 1
    row = real_conn.execute("SELECT value, value_status FROM visitor_stats").fetchone()
    assert (row["value"], row["value_status"]) == (120, "final")
    revs = {r["field"]: (r["old_value"], r["new_value"]) for r in real_conn.execute("SELECT * FROM revisions")}
    assert revs["value"] == ("100.0", "120.0")


def test_conflicting_rows_in_same_file_reject_import(real_conn):
    csv = (VS_HEADER + "2025-01,Japan,United States,visitor_arrivals,100,persons,final,JNTO,2025-02-01\n"
           + "2025-01,Japan,United States,visitor_arrivals,999,persons,final,JNTO,2025-02-01\n")
    batch = run_import(real_conn, "real", "visitor_stats", "c.csv", csv.encode())
    assert batch["status"] == "rejected"
    assert count(real_conn, "visitor_stats") == 0


def test_japan_and_tokyo_are_different_records(real_conn):
    csv = (VS_HEADER + "2025-01,Japan,United States,visitor_arrivals,100,persons,final,JNTO,2025-02-01\n"
           + "2025-01,Tokyo,United States,visitor_arrivals,40,persons,final,TMG,2025-02-01\n")
    batch = run_import(real_conn, "real", "visitor_stats", "d.csv", csv.encode())
    assert batch["rows_inserted"] == 2


def test_blank_competitor_price_is_missing_and_price_requires_currency(real_conn):
    header = "business,tour_name,price,currency,language,date_observed\n"
    ok = run_import(real_conn, "real", "competitor_offers", "o.csv",
                    (header + "A,Walk,,,English,2025-01-01\n").encode())
    assert ok["status"] == "success"
    assert real_conn.execute("SELECT price FROM competitor_offers").fetchone()[0] is None
    bad = run_import(real_conn, "real", "competitor_offers", "o2.csv",
                     (header + "B,Walk,5000,,English,2025-01-01\n").encode())
    assert bad["status"] == "rejected" and "currency" in bad["errors"][0]["message"]


def test_original_file_is_preserved(real_conn, isolated_data_dir):
    content = (FEEDBACK_HEADER + "Lovely walk,en,Survey,5,Own survey,,2025-02-01\n").encode()
    batch = run_import(real_conn, "real", "feedback", "my feedback.csv", content)
    from pathlib import Path
    assert Path(batch["raw_path"]).read_bytes() == content
    assert str(isolated_data_dir) in batch["raw_path"]


def test_unsupported_file_type_rejected(real_conn):
    batch = run_import(real_conn, "real", "feedback", "notes.pdf", b"%PDF")
    assert batch["status"] == "rejected" and "Unsupported" in batch["errors"][0]["message"]


def _jnto_workbook(all_final: bool) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2026"
    ws["A1"] = "2026年　訪日外客数（総数）"
    ws.append([])
    ws.append([])
    ws.append([None, None] + [x for m in range(1, 13) for x in (f"{m}月", "伸率")])  # row 4
    ws.append(["総数", None, 1000, 1.0, 1100, 2.0, 1200, 3.0])  # Apr+ blank = unpublished
    ws.append(["米国", None, 200, 1.0, None, None, 0, 0.0])   # Feb unpublished, Mar genuinely 0
    ws.cell(row=5, column=7).font = Font(italic=True)          # G5 = total Mar, an estimate
    ws.append(["注２：　" + ("表中の数値は全て確定値である。" if all_final else "斜体の数値は推計値である。")])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_jnto_workbook_blank_cells_are_missing_and_italics_are_estimates(real_conn):
    batch = run_import(real_conn, "real", "visitor_stats", "20260916_1615-5.xlsx", _jnto_workbook(False),
                       parser=parse_jnto_monthly_workbook)
    assert batch["status"] == "success"
    data = {(r["visitor_origin"], r["reporting_month"]): r for r in real_conn.execute("SELECT * FROM visitor_stats")}
    assert ("United States", "2026-02") not in data            # blank -> no record
    assert data[("United States", "2026-03")]["value"] == 0     # an explicit 0 stays 0
    assert data[("All origins", "2026-03")]["value_status"] == "estimate"
    assert data[("All origins", "2026-01")]["value_status"] == "provisional"
    assert data[("All origins", "2026-01")]["geography"] == "Japan"
    assert data[("All origins", "2026-01")]["publication_date"] == "2026-09-16"
    assert len(data) == 5


def test_jnto_final_sheet_marks_values_final(real_conn):
    run_import(real_conn, "real", "visitor_stats", "x.xlsx", _jnto_workbook(True), parser=parse_jnto_monthly_workbook)
    assert {r[0] for r in real_conn.execute("SELECT DISTINCT value_status FROM visitor_stats")} == {"final"}


def test_wrong_workbook_is_rejected(real_conn):
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    buf = io.BytesIO()
    wb.save(buf)
    batch = run_import(real_conn, "real", "visitor_stats", "x.xlsx", buf.getvalue(), parser=parse_jnto_monthly_workbook)
    assert batch["status"] == "rejected"
