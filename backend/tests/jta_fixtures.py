"""Small workbooks laid out like the Japan Tourism Agency files (same sheet names,
header rows and column positions as the 2026-Q2 releases), for tests."""
import io

import openpyxl


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def national_workbook(period="2026年4-6月期 【2次速報】", tour_spend=2226.7, tour_rate=8.7562,
                      tour_buyers=676, korea_tour_rate=None) -> bytes:
    """参考2 (spend per visitor) and 表2-1 (purchase rate, spend per buyer) for two segments."""
    wb = openpyxl.Workbook()
    wb.active.title = "表紙"
    wb.active["B2"] = "インバウンド消費動向調査"

    ref = wb.create_sheet("参考2")
    ref["B2"] = "参考2　国籍･地域（24区分）別　費目別１人１回当たり旅行消費単価 （パッケージツアー参加費内訳含む）"
    ref["B4"] = period
    ref["B5"], ref["E5"], ref["G5"] = "調査項目", "全国籍･地域", "韓国"
    ref["E6"], ref["F6"], ref["G6"], ref["H6"] = "消費単価", "構成比", "消費単価", "構成比"
    ref.append([None, "全体　【A1】", None, None, 244733.2571, 100, 99255.8542, 100])
    ref.append([None, "日本滞在中の", "宿泊費", None, 89787.0323, 36.7, 30874.7, 31.1])
    ref.append([None, "【A1】", "娯楽等サービス費", None, 10620.0784, 4.3, 3973.1, 4.0])
    ref.append([None, None, None, "現地ツアー・観光ガイド", tour_spend, 0.9, 582.3, 0.6])
    ref.append([None, None, None, "温泉・温浴施設・ｴｽﾃ・ﾘﾗｸｾﾞｰｼｮﾝ", 350.75, 0.14, None, None])  # blank = not published
    ref.append([None, None, None, "謎の新項目", 10.0, 0.0, 1.0, 0.0])  # unknown label is kept, with a warning
    ref.append([None, "注）サンプル", None, None, 1, 1, 1, 1])  # footnote ends the block

    t21 = wb.create_sheet("表2-1")
    t21["B2"] = "表2-1　国籍･地域（24区分）別　費目別購入率および購入者単価"
    t21["B4"] = period
    t21["B5"], t21["E5"], t21["G5"] = "調査項目", "全国籍･地域", "韓国"
    t21["E6"], t21["F6"], t21["G6"], t21["H6"] = "回答数", "購入率", "回答数", "購入率"
    t21.append([None, "旅行前支出", "団体パッケージツアー", None, 694, 8.7789, 56, 5.9434])
    t21.append([None, "【A1】", "娯楽等サービス費", None, 4562, 56.0227, 358, 44.467])
    t21.append([None, None, None, "現地ツアー・観光ガイド", tour_buyers, tour_rate, 77, korea_tour_rate])
    t21.append([])
    t21.append([None, "調査項目", None, None, "全国籍･地域", None, "韓国"])
    t21.append([None, None, None, None, "回答数", "購入者単価", "回答数", "購入者単価"])
    t21.append([None, "【A1】", "娯楽等サービス費", None, 4562, 16714.67, 358, 8283.44])
    t21.append([None, None, None, "現地ツアー・観光ガイド", tour_buyers, 21995.65, 77, 5521.91])
    return _save(wb)


def prefecture_workbook(period="2026年（令和8年）4-6月期", tokyo_total=9426.87, tokyo_entertainment=342.13) -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = "目次"
    wb.active["A1"] = "インバウンド消費動向調査"

    t11 = wb.create_sheet("表1-1")
    t11["A3"], t11["B3"] = "表1-1", "都道府県（47区分）別　訪問者数および消費単価　【全目的】"
    t11["A5"], t11["E5"], t11["G5"] = period, "（単位：万人）", "（単位：万円／人）"
    t11["D6"], t11["E6"], t11["G6"] = "訪問率", "訪問者数", "消費単価注1"
    t11["C7"], t11["F7"] = "標本サイズ（人）", "標本サイズ（人）"
    t11["B8"] = "訪問地"
    t11.append([13, "東京都", 13130, 0.5281877, 541.6161770, 7622, 17.4050675])
    t11.append([14, "神奈川県", 2000, 0.1, 100.0, 1500, 5.0])

    t13 = wb.create_sheet("表1-3")
    t13["A3"], t13["B3"] = "表1-3", "都道府県（47区分）別，費目（7区分）別　旅行消費額　【全目的】"
    t13["A5"], t13["J5"] = period, "（単位：億円）"
    t13["C6"], t13["D6"] = "旅行\n消費額注1", "費目別（7区分）"
    for col, text in zip("DEFGHIJ", ["団体･パック", "宿泊費", "飲食費", "交通費", "娯楽等", "買物代", "その他"]):
        t13[f"{col}7"] = text
    t13["B8"], t13["D8"], t13["H8"] = "訪問地", "参加費", "サービス費"
    t13.append([13, "東京都", tokyo_total, 614.91, 3386.15, 1842.64, 243.55, tokyo_entertainment, 2994.67, 2.81])
    t13.append([None, "三大都市圏注2", 17616.8, 1410.0, 6020.3, 3545.6, 435.7, 773.0, 5420.4, 11.8])  # region: skipped
    return _save(wb)


JNTO_PAGE_HTML = """<html><body>
<a href="/statistics/data/_files/20260916_1615-5.xlsx">国籍/月別 訪日外客数（2003年～2026年）</a>
<a href="/statistics/data/_files/other.pdf">報道発表資料</a>
</body></html>"""

JTA_PAGE_HTML = """<html><body>
<h2>最新の調査結果を見る</h2>
<p>最新の調査結果 &nbsp; 2026年4～6月期 2次速報（公表日：2026年9月30日）</p>
<a href="/kankocho/content/LATEST.xlsx">➤集計表</a>
<a href="#">これまでの調査結果を見る&darr;</a>
<h2>これまでの調査結果を見る</h2>
<details><summary>2026年</summary>
  <details><summary>　調査概要</summary><ul><li><a href="/kankocho/content/S1.pdf">4-6月期(2次速報)</a></li></ul></details>
  <details><summary>　集計表</summary><ul><li><a href="/kankocho/content/N2026Q2.xlsx">4-6月期(2次速報)</a></li></ul></details>
  <details><summary>　【参考】都道府県別集計表</summary><ul><li><a href="/kankocho/content/P2026Q2.xlsx">4-6月期</a></li></ul></details>
</details>
<details><summary>2024年　※1-3月期以前は旧調査</summary>
  <details><summary>　集計表</summary><ul>
    <li><a href="/kankocho/content/Y2024.xls">暦年</a></li>
    <li><a href="/kankocho/content/N2024Q1.xls">1-3月期</a></li>
  </ul></details>
</details>
</body></html>"""
