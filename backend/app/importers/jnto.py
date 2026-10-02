"""Adapter for JNTO's monthly 'visitor arrivals by nationality/region' workbook.

Source page: https://www.jnto.go.jp/statistics/data/visitors-statistics/
Download the XLSX manually and import it here. We do not scrape the site.

Workbook layout (checked against the 2003-2026 file published 2026-09-16):
- one sheet per year ('2026', '2025', ...)
- a header row containing '1月', '伸率', '2月', '伸率', ... '累計' (cumulative)
- one row per market; sub-markets (e.g. Israel under Middle East) sit in the
  second label column
- a blank cell = not yet published -> we create NO record (missing, not zero)
- footnote '注２' says whether all figures are final (確定値); otherwise italic
  cells are estimates (推計値) and the rest are treated as provisional

We ignore JNTO's own growth-rate (伸率) columns and recompute changes
ourselves, only where both periods exist.

IMPORTANT: these are arrivals to JAPAN, not visitors to Tokyo, and the
breakdown is by nationality, which says nothing about tour-language preference.
"""
import io
import re
import unicodedata
from datetime import date
from pathlib import Path

import openpyxl

from app.importers.common import ParseError, parse_date
from app.importers.service import ParsedFile, RowError

# Japanese label -> (English name, origin_level)
LABELS: dict[str, tuple[str, str]] = {
    "総数": ("All origins", "total"),
    "アジア計": ("Asia (total)", "region"),
    "韓国": ("South Korea", "country"),
    "中国": ("China", "country"),
    "台湾": ("Taiwan", "country"),
    "香港": ("Hong Kong", "country"),
    "タイ": ("Thailand", "country"),
    "シンガポール": ("Singapore", "country"),
    "マレーシア": ("Malaysia", "country"),
    "インドネシア": ("Indonesia", "country"),
    "フィリピン": ("Philippines", "country"),
    "ベトナム": ("Vietnam", "country"),
    "インド": ("India", "country"),
    "中東地域": ("Middle East (total)", "region"),
    "イスラエル": ("Israel", "country"),
    "トルコ": ("Türkiye", "country"),
    "GCC6か国": ("GCC 6 countries", "region"),
    "マカオ": ("Macau", "country"),
    "モンゴル": ("Mongolia", "country"),
    "その他アジア": ("Other Asia", "other"),
    "ヨーロッパ計": ("Europe (total)", "region"),
    "英国": ("United Kingdom", "country"),
    "フランス": ("France", "country"),
    "ドイツ": ("Germany", "country"),
    "イタリア": ("Italy", "country"),
    "スペイン": ("Spain", "country"),
    "ロシア": ("Russia", "country"),
    "北欧地域": ("Nordic countries (total)", "region"),
    "スウェーデン": ("Sweden", "country"),
    "デンマーク": ("Denmark", "country"),
    "ノルウェー": ("Norway", "country"),
    "フィンランド": ("Finland", "country"),
    "オランダ": ("Netherlands", "country"),
    "スイス": ("Switzerland", "country"),
    "ベルギー": ("Belgium", "country"),
    "ポーランド": ("Poland", "country"),
    "オーストリア": ("Austria", "country"),
    "ポルトガル": ("Portugal", "country"),
    "アイルランド": ("Ireland", "country"),
    "その他ヨーロッパ": ("Other Europe", "other"),
    "アフリカ計": ("Africa (total)", "region"),
    "北アメリカ計": ("North America (total)", "region"),
    "米国": ("United States", "country"),
    "カナダ": ("Canada", "country"),
    "メキシコ": ("Mexico", "country"),
    "その他北アメリカ": ("Other North America", "other"),
    "南アメリカ計": ("South America (total)", "region"),
    "ブラジル": ("Brazil", "country"),
    "その他南アメリカ": ("Other South America", "other"),
    "オセアニア計": ("Oceania (total)", "region"),
    "豪州": ("Australia", "country"),
    "ニュージーランド": ("New Zealand", "country"),
    "その他オセアニア": ("Other Oceania", "other"),
    "無国籍・その他": ("Stateless / other", "other"),
}

MONTH_HEADER = re.compile(r"^(\d{1,2})月$")


def _norm(value) -> str:
    return unicodedata.normalize("NFKC", str(value)).strip() if value is not None else ""


def publication_date_from_filename(filename: str) -> str | None:
    """JNTO file names start with the publication date, e.g. 20260916_1615-5.xlsx."""
    m = re.match(r"(\d{4})(\d{2})(\d{2})_", Path(filename).name)
    if not m:
        return None
    try:
        return parse_date(f"{m.group(1)}-{m.group(2)}-{m.group(3)}")
    except ParseError:
        return None


def parse_jnto_monthly_workbook(content: bytes, filename: str, collection_date: str | None = None) -> ParsedFile:
    result = ParsedFile(importer="jnto_monthly_xlsx")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        result.errors.append(RowError(None, None, f"Could not open the workbook: {exc}"))
        return result

    publication_date = publication_date_from_filename(filename)
    if publication_date:
        result.warnings.append(f"Publication date {publication_date} was inferred from the file name.")
    else:
        result.warnings.append("Publication date unknown (file name does not start with YYYYMMDD_).")
    collection_date = collection_date or date.today().isoformat()

    year_sheets = [ws for ws in wb.worksheets if re.fullmatch(r"\d{4}", ws.title.strip())]
    if not year_sheets:
        result.errors.append(RowError(None, None,
                                      "No year sheets (e.g. '2025') found. Is this the JNTO monthly "
                                      "arrivals-by-nationality workbook?"))
        return result

    unknown_labels: set[str] = set()
    for ws in year_sheets:
        year = int(ws.title.strip())
        all_final = any("確定値" in _norm(c.value) and ("全て" in _norm(c.value) or "すべて" in _norm(c.value))
                        for row in ws.iter_rows() for c in row if isinstance(c.value, str))

        header_row, month_cols = None, {}
        for row in ws.iter_rows():
            found = {c.column: int(MONTH_HEADER.match(_norm(c.value)).group(1))
                     for c in row if MONTH_HEADER.match(_norm(c.value))}
            if len(found) >= 6:
                header_row, month_cols = row[0].row, found
                break
        if header_row is None:
            result.errors.append(RowError(None, None, f"Sheet {ws.title}: could not find the month header row (1月, 2月, ...)."))
            continue
        first_value_col = min(month_cols)

        for row in ws.iter_rows(min_row=header_row + 1):
            label_cells = [c for c in row if c.column < first_value_col and _norm(c.value)]
            if not label_cells:
                continue
            label = _norm(label_cells[-1].value)
            if label.startswith(("注", "※", "出典")):
                break  # footnotes
            english, level = LABELS.get(label, (label, "other"))
            if label not in LABELS:
                unknown_labels.add(label)

            for col, month in month_cols.items():
                cell = ws.cell(row=row[0].row, column=col)
                if cell.value is None or _norm(cell.value) in ("", "-", "‐", "―"):
                    continue  # not published: missing, not zero
                if not isinstance(cell.value, (int, float)):
                    result.errors.append(RowError(row[0].row, f"{ws.title}!{cell.coordinate}",
                                                  f"Expected a number for {label} {year}-{month:02d}, got '{cell.value}'"))
                    continue
                status = "final" if all_final else ("estimate" if cell.font and cell.font.i else "provisional")
                result.records.append({
                    "reporting_month": f"{year:04d}-{month:02d}",
                    "geography": "Japan",
                    "visitor_origin": english,
                    "origin_level": level,
                    "metric": "visitor_arrivals",
                    "value": float(cell.value),
                    "unit": "persons",
                    "value_status": status,
                    "source": "JNTO",
                    "original_label": label,
                    "publication_date": publication_date,
                    "collection_date": collection_date,
                    "_row": row[0].row,
                })

    if unknown_labels:
        result.warnings.append("Labels without an English translation were kept as-is: "
                               + ", ".join(sorted(unknown_labels)))
    if any(r["value_status"] != "final" for r in result.records):
        result.warnings.append("Some figures are provisional or estimates (italic in the workbook). "
                               "Re-import a later JNTO file to update them; old values are kept in revision history.")
    return result
